#!/usr/bin/env python3
"""Build Sunset Sage ↔ Furniture Information Sheet crosswalk."""

from __future__ import annotations

import csv
from pathlib import Path

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
XLS = ROOT / "docs" / "Furniture Information Sheet.xls"
INV = ROOT / "docs" / "sunset-sage" / "inventory.csv"
OUT_CSV = ROOT / "docs" / "sunset-sage" / "crosswalk.csv"
OUT_XLSX = ROOT / "docs" / "sunset-sage" / "crosswalk.xlsx"

# ss_row -> (sku, confidence, notes)
# high = name+hide agree; medium = best SKU, confirm; review = photo/SKU gap; new = not on master
M = {
    2: ("55343454", "high", "Vegas Short — hourglass stool"),
    3: ("30400123", "medium", "Generic brown mix; mapped to Saddle Brown Hide. Confirm."),
    4: ("55344545", "high", "Vegas Tall — hourglass bar height"),
    5: ("30400238", "high", "Ivory Leaf (sheet said floral)"),
    6: ("30400002", "high", "Catalog hide is Brown Croc"),
    7: ("30450001", "medium", "Generic Smith; mapped to Tufted Brown"),
    8: ("30102003", "medium", "Trophy 2 White Croc — not the tufted corset"),
    9: ("10220789", "high", ""),
    10: ("30810276", "high", "Brown diamond leather; hide back on photo"),
    11: ("30217165", "high", "New Cow Brindle Hide"),
    12: ("30400123", "medium", "Cowhide; also used for generic brown mix row 3"),
    13: ("30400099", "high", ""),
    14: ("10220790", "high", ""),
    15: ("30400725", "high", ""),
    16: ("30400234", "high", ""),
    17: ("11646222", "high", "Serena Dark"),
    18: ("11507002", "high", "Lynn Brown croc"),
    19: ("30400137", "high", ""),
    20: ("30810277", "medium", "Closest Kennedy: brown leather, hide back (not named croc)"),
    21: ("30400240", "high", "Copper Rose"),
    22: ("50213087", "high", "Brown, Diamond, Flower Stitch"),
    23: ("30103002", "high", "Trophy 3 Brown Croc"),
    24: ("50750001", "high", ""),
    25: ("", "new", "Scott line not on master sheet"),
    26: ("", "new", "Scott line not on master sheet"),
    27: ("", "new", "Scott line not on master sheet"),
    28: ("", "review", "Smith tufted white croc — catalog photo SKU 30190002 is not on the master"),
    29: ("30400237", "high", "Turquoise Leaf (not Floral)"),
    30: ("30103016", "high", ""),
    31: ("", "new", "Dakota only has Black Hide on the master; this is brown hide"),
    32: ("30400724", "high", ""),
    33: ("30400003", "high", ""),
    34: ("10836088", "high", "Only Dakota SKU on the master"),
    35: ("30102013", "high", "White Croc Weave Brown"),
    36: ("10220791", "high", "White, Black & White Hide"),
    37: ("30311002", "high", ""),
    38: ("30400236", "high", ""),
    39: ("30450003", "high", ""),
    40: ("30400241", "high", "Laredo Brown"),
    41: ("30810640", "high", ""),
    42: ("30810638", "high", ""),
    43: ("", "new", "Saddle Black Croc not on master"),
    44: ("10220791", "medium", "Tri-color with nailhead — confirm vs salt-and-pepper white"),
    45: ("", "new", "Arc chair not on master"),
    46: ("11308002", "high", "Grace 2 Brown Croc"),
    47: ("30110048", "high", ""),
    48: ("30400242", "high", "Saddle Top Blue (not Blue Croc)"),
    49: ("30103050", "high", ""),
    50: ("11307130", "high", "Grace Axis — price $1000 matches"),
    51: ("30103015", "medium", "Trophy 3 Tufted Medium Brown — mocha"),
    52: ("10220791", "medium", "Tri-color no nailhead — confirm vs row 44"),
    53: ("11207007", "high", "Croc & Damask / Baroque"),
    54: ("50538003", "high", ""),
    55: ("30311004", "high", ""),
    56: ("", "new", "Chelsea not on master"),
    57: ("50405101", "high", "Horn Brown Croc"),
    58: ("30311005", "high", ""),
    59: ("30400235", "high", ""),
    60: ("30450724", "high", ""),
    61: ("", "new", "Horn tufted brown not on master"),
    62: ("50320130", "high", "Pearl Axis / Springbuck hide"),
    63: ("30400769", "high", ""),
    64: ("30310123", "medium", "Bethany Brown Hide — no-arms not explicit on master"),
    65: ("", "new", "Bethany with arms Loredo — master Loredo is no-arms"),
    66: ("50412789", "high", ""),
    67: ("", "new", "Brooke barstool (not dining)"),
    68: ("30103002", "medium", "Navy / dark croc full front — confirm vs Trophy 3 Brown Croc"),
    69: ("", "new", "Bethany with arms Springbuck not on master"),
    70: ("11111", "medium", "Placeholder SKU Brooke Brown Barstool"),
    71: ("50405100", "high", ""),
    72: ("11308002", "high", "Same SKU as row 46 (two qty lots)"),
    73: ("50307100", "medium", "Grace Office Chair White Yolk — armless"),
    74: ("30102014", "high", "Trophy 2 Light Caramel"),
    75: ("50306101", "high", "Grace Arm Tufted Office Chair"),
    76: ("30103002", "high", "Same Trophy 3 Brown Croc as row 23"),
    77: ("11307637", "high", ""),
    78: ("", "new", "Jack cowhide + turquoise trim not on master"),
    79: ("", "new", "Half-moon ottoman not on master"),
    80: ("50320130", "medium", "Pearl cowhide — may be same as Axis/Springbuck"),
    81: ("", "new", "Texas King Bed not on master"),
    82: ("50370267", "medium", "Waller generic SKU (empty hide on master)"),
    83: ("50701836", "high", "Brown Croc front, Dark Brindle hide back"),
    84: ("30310123", "medium", "Bethany Brown Hide with arms — confirm vs no-arms row 64"),
    85: ("11507005", "high", ""),
    86: ("55356430", "medium", "Murphy Round 37\" — size not on SS row"),
    87: ("50763099", "high", "Britney Diamond Weave"),
    88: ("55356431", "medium", "Murphy Round 43\" taupe — size assumed"),
    89: ("50701099", "high", ""),
    90: ("30810088", "high", ""),
    91: ("11307973", "high", ""),
    92: ("55356420", "medium", "Murphy Springbuck / taupe ottoman"),
    93: ("50320003", "high", ""),
    94: ("55355430", "high", "Jack Rectangle"),
    95: ("50784222", "high", "Only Hugh on master"),
    96: ("50701835", "high", "Brindle hide front, Brown Croc back"),
    97: ("50412790", "high", ""),
    98: ("55355431", "high", "Jack Taupe Cowboy Tool"),
    99: ("", "new", "Dakota Taupe not on master"),
}


def sku_key(raw) -> str:
    s = str(raw).strip()
    try:
        if float(s) == int(float(s)):
            return str(int(float(s)))
    except Exception:
        pass
    return s


def money(v):
    try:
        return float(v)
    except Exception:
        return None


def qty(v):
    try:
        return float(str(v).split()[0])
    except Exception:
        return 0.0


def load_catalog():
    wb = xlrd.open_workbook(XLS)
    sh = wb.sheet_by_name("Catalog")
    headers = [sh.cell_value(0, c) for c in range(sh.ncols)]
    rows = []
    for r in range(1, sh.nrows):
        d = {headers[c]: sh.cell_value(r, c) for c in range(sh.ncols)}
        d["sku"] = sku_key(d["sku"])
        rows.append(d)
    return rows


def main() -> None:
    catalog = load_catalog()
    by_sku = {c["sku"]: c for c in catalog}
    ss = list(csv.DictReader(INV.open()))

    out_rows = []
    for r in ss:
        rowno = int(r["row"])
        sku, conf, notes = M[rowno]
        cat = by_sku.get(sku, {})
        ss_price = money(r["price"])
        cat_price = money(cat.get("price")) if cat else None
        if ss_price is not None and cat_price is not None and ss_price != cat_price:
            delta = f"{cat_price:.0f} on master vs {ss_price:.0f} at Sunset Sage"
            notes = f"{notes} | {delta}" if notes else delta
        out_rows.append(
            {
                "ss_row": rowno,
                "ss_item": r["item"],
                "ss_qty": r["quantity"],
                "ss_price": r["price"],
                "ss_photo": r["photo"],
                "match": conf,
                "sku": sku,
                "catalog_name": cat.get("name", ""),
                "catalog_hide": cat.get("hide", ""),
                "catalog_price": cat.get("price", ""),
                "catalog_category": cat.get("category", ""),
                "notes": notes,
            }
        )

    hit = {r["sku"] for r in out_rows if r["sku"]}
    unmatched_catalog = [c for c in catalog if c["sku"] not in hit]

    fields = [
        "ss_row",
        "ss_item",
        "ss_qty",
        "ss_price",
        "ss_photo",
        "match",
        "sku",
        "catalog_name",
        "catalog_hide",
        "catalog_price",
        "catalog_category",
        "notes",
    ]
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(out_rows)

    yellow = PatternFill("solid", fgColor="F6E7B2")
    green = PatternFill("solid", fgColor="D9E8D3")
    orange = PatternFill("solid", fgColor="F7D7B5")
    blue = PatternFill("solid", fgColor="D6E3F0")
    header_fill = PatternFill("solid", fgColor="2C2416")
    header_font = Font(bold=True, color="F4EFE6")
    fills = {"high": green, "medium": orange, "review": yellow, "new": blue}

    wb = Workbook()
    ws = wb.active
    ws.title = "Crosswalk"
    ws.freeze_panes = "A2"
    head = [
        "SS row",
        "Sunset Sage item",
        "On hand",
        "SS price",
        "Photo",
        "Match",
        "SKU",
        "Master name",
        "Master hide",
        "Master price",
        "Category",
        "Notes",
    ]
    for c, h in enumerate(head, 1):
        cell = ws.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
    for i, r in enumerate(out_rows, 2):
        vals = [
            r["ss_row"],
            r["ss_item"],
            qty(r["ss_qty"]),
            money(r["ss_price"]),
            r["ss_photo"],
            r["match"],
            r["sku"],
            r["catalog_name"],
            r["catalog_hide"],
            money(r["catalog_price"]) if r["catalog_price"] != "" else "",
            r["catalog_category"],
            r["notes"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            if c == 6:
                cell.fill = fills.get(r["match"], yellow)
    ws.auto_filter.ref = f"A1:L{len(out_rows)+1}"
    for i, w in enumerate([10, 48, 12, 12, 42, 12, 14, 16, 42, 14, 18, 55], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "Sunset Sage × Furniture Information Sheet"
    sm["A1"].font = Font(size=18, bold=True)
    counts = {}
    for r in out_rows:
        counts[r["match"]] = counts.get(r["match"], 0) + 1
    sm["A3"] = "Match"
    sm["B3"] = "Rows"
    sm["C3"] = "On-hand units"
    for col in range(1, 4):
        sm.cell(3, col).fill = header_fill
        sm.cell(3, col).font = header_font
    for i, m in enumerate(["high", "medium", "review", "new"], 4):
        sm.cell(i, 1, m).fill = fills[m]
        sm.cell(i, 2, counts.get(m, 0))
        sm.cell(i, 3, sum(qty(r["ss_qty"]) for r in out_rows if r["match"] == m))
    sm["A9"] = (
        "high = name + hide agree. medium = best SKU, confirm hide. "
        "review = photo exists but SKU missing from master. new = not on the master sheet."
    )
    sm["A11"] = "Master SKUs with no Sunset Sage row"
    sm["A11"].font = Font(bold=True)
    sm["A12"] = len(unmatched_catalog)
    sm["B12"] = "of"
    sm["C12"] = len(catalog)
    sm["A14"] = "Price differences are in Notes on the Crosswalk tab."
    for col, w in enumerate([28, 12, 18, 14], 1):
        sm.column_dimensions[get_column_letter(col)].width = w

    nw = wb.create_sheet("New at Sunset Sage")
    for c, h in enumerate(["SS item", "On hand", "SS price", "Photo", "Why"], 1):
        cell = nw.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
    i = 2
    for r in out_rows:
        if r["match"] in ("new", "review"):
            nw.cell(i, 1, r["ss_item"])
            nw.cell(i, 2, qty(r["ss_qty"]))
            nw.cell(i, 3, money(r["ss_price"]))
            nw.cell(i, 4, r["ss_photo"])
            nw.cell(i, 5, r["notes"])
            i += 1
    for col, w in enumerate([48, 12, 12, 42, 70], 1):
        nw.column_dimensions[get_column_letter(col)].width = w

    um = wb.create_sheet("On master, not in SS")
    for c, h in enumerate(["SKU", "Name", "Hide", "Price", "Category"], 1):
        cell = um.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
    for i, c in enumerate(unmatched_catalog, 2):
        um.cell(i, 1, c["sku"])
        um.cell(i, 2, c["name"])
        um.cell(i, 3, c["hide"])
        um.cell(i, 4, c["price"])
        um.cell(i, 5, c["category"])
    um.auto_filter.ref = f"A1:E{len(unmatched_catalog)+1}"
    for col, w in enumerate([14, 16, 42, 12, 20], 1):
        um.column_dimensions[get_column_letter(col)].width = w

    wb.save(OUT_XLSX)
    print("crosswalk rows", len(out_rows), counts)
    print("unmatched master SKUs", len(unmatched_catalog))
    print("wrote", OUT_CSV.name, OUT_XLSX.name)


if __name__ == "__main__":
    main()
